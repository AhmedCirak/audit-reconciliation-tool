"""Export GL, Bank, Reconciliation i Summary u jedan Excel workbook."""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

OUT = "../output/Audit_Reconciliation.xlsx"

rec = pd.read_pickle("./_rec.pkl")
gl = pd.read_pickle("./_gl.pkl")
bank = pd.read_pickle("./_bank.pkl")
stats = pd.read_pickle("./_stats.pkl")

# Yes/No umjesto TRUE/FALSE -> lakse za COUNTIF, pivot i conditional formatting
for df in (rec,):
    for c in ("Is_Exception", "Has_Duplicate_ID"):
        df[c] = df[c].map({True: "Yes", False: "No"})
for df in (gl, bank):
    df["Is_Duplicate"] = df["Is_Duplicate"].map({True: "Yes", False: "No"})

# ---------------------------------------------------------------- styles
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
TITLE = Font(name=FONT, bold=True, size=16, color="1F3864")
SUB = Font(name=FONT, italic=True, size=9, color="595959")
SEC = Font(name=FONT, bold=True, size=11, color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CUR = '#,##0.00;[Red](#,##0.00);"-"'
DATE_FMT = "yyyy-mm-dd"
MON_FMT = "yyyy-mm"

wb = Workbook()
wb.remove(wb.active)


def write_table(ws, df, table_name, widths=None, number_formats=None):
    """Upisi DataFrame i registruj ga kao pravu Excel Table."""
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    for row in df.itertuples(index=False):
        ws.append(list(row))

    n_rows, n_cols = len(df) + 1, len(df.columns)
    ref = f"A1:{get_column_letter(n_cols)}{n_rows}"

    for row in ws.iter_rows(min_row=2, max_row=n_rows, max_col=n_cols):
        for cell in row:
            cell.font = BODY

    if number_formats:
        for col_name, fmt in number_formats.items():
            idx = list(df.columns).index(col_name) + 1
            for r in range(2, n_rows + 1):
                ws.cell(row=r, column=idx).number_format = fmt

    tbl = Table(displayName=table_name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(tbl)

    for i, col_name in enumerate(df.columns, start=1):
        w = (widths or {}).get(col_name, max(11, min(24, len(str(col_name)) + 4)))
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    return n_rows


# ============================================================ GL_Data
ws_gl = wb.create_sheet("GL_Data")
write_table(
    ws_gl, gl, "tblGL",
    widths={"Description": 28, "Transaction_ID": 15, "Category": 22, "Date": 12},
    number_formats={"Amount": CUR, "Date": DATE_FMT},
)

# ============================================================ Bank_Data
ws_bk = wb.create_sheet("Bank_Data")
write_table(
    ws_bk, bank, "tblBank",
    widths={"Description": 28, "Transaction_ID": 15, "Date": 12},
    number_formats={"Amount": CUR, "Date": DATE_FMT},
)

# ============================================================ Reconciliation
ws_rc = wb.create_sheet("Reconciliation")
n_rec = write_table(
    ws_rc, rec, "tblRecon",
    widths={"Transaction_ID": 15, "Category": 24, "Description_GL": 26,
            "Description_Bank": 26, "Status": 17, "Is_Exception": 13,
            "Has_Duplicate_ID": 17, "Date": 12, "Month": 11},
    number_formats={"Amount_GL": CUR, "Amount_Bank": CUR, "Amount": CUR,
                    "Abs_Amount": CUR, "Amount_Diff": CUR,
                    "Date": DATE_FMT, "Month": MON_FMT},
)

COLS = {c: get_column_letter(i) for i, c in enumerate(rec.columns, start=1)}
C_STAT, C_AMT, C_DUP = COLS["Status"], COLS["Amount"], COLS["Has_Duplicate_ID"]
R_STAT = f"Reconciliation!${C_STAT}$2:${C_STAT}${n_rec}"
R_AMT = f"Reconciliation!${C_AMT}$2:${C_AMT}${n_rec}"
R_DUP = f"Reconciliation!${C_DUP}$2:${C_DUP}${n_rec}"
R_ID = f"Reconciliation!$A$2:$A${n_rec}"

# ============================================================ Summary
ws = wb.create_sheet("Summary")
ws.sheet_view.showGridLines = False

ws["B2"] = "Audit Reconciliation — Summary"
ws["B2"].font = TITLE
ws["B3"] = ("General Ledger vs Bank Statement · sve KPI vrijednosti su formule "
            "vezane na sheet 'Reconciliation' i azuriraju se automatski")
ws["B3"].font = SUB

KPI_FILL = PatternFill("solid", fgColor="F2F5FA")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
BAD_FILL = PatternFill("solid", fgColor="FCE4E4")

kpis = [
    ("Total Transactions", f'=COUNTA({R_ID})', "neutral"),
    ("Matched", f'=COUNTIF({R_STAT},"Matched")', "good"),
    ("Exceptions", f'=COUNTA({R_ID})-COUNTIF({R_STAT},"Matched")', "bad"),
    ("Duplicates", f'=COUNTIF({R_DUP},"Yes")', "bad"),
    ("Only in GL", f'=COUNTIF({R_STAT},"Only in GL")', "bad"),
    ("Only in Bank", f'=COUNTIF({R_STAT},"Only in Bank")', "bad"),
    ("Amount Mismatch", f'=COUNTIF({R_STAT},"Amount Mismatch")', "bad"),
]

# KPI kartice: 4 u prvom redu, 3 u drugom
start_row, per_row = 5, 4
for i, (label, formula, kind) in enumerate(kpis):
    r = start_row + (i // per_row) * 4
    c = 2 + (i % per_row) * 2
    lab, val = ws.cell(row=r, column=c), ws.cell(row=r + 1, column=c)
    lab.value, val.value = label, formula
    lab.font = Font(name=FONT, bold=True, size=9, color="595959")
    lab.alignment = Alignment(horizontal="center")
    val.font = Font(name=FONT, bold=True, size=22,
                    color={"good": "1E6B34", "bad": "9C1C1C"}.get(kind, "1F3864"))
    val.alignment = Alignment(horizontal="center")
    val.number_format = "#,##0"
    fill = {"good": OK_FILL, "bad": BAD_FILL}.get(kind, KPI_FILL)
    for cell in (lab, val):
        cell.fill, cell.border = fill, BOX
    ws.row_dimensions[r + 1].height = 32

# Match rate
r_mr = start_row + 8
ws.cell(row=r_mr, column=2, value="Match Rate").font = Font(
    name=FONT, bold=True, size=9, color="595959")
mr = ws.cell(row=r_mr + 1, column=2,
             value=f'=IFERROR(COUNTIF({R_STAT},"Matched")/COUNTA({R_ID}),0)')
mr.font = Font(name=FONT, bold=True, size=22, color="1E6B34")
mr.number_format = "0.0%"
mr.alignment = Alignment(horizontal="center")
ws.cell(row=r_mr, column=2).alignment = Alignment(horizontal="center")
for cell in (ws.cell(row=r_mr, column=2), mr):
    cell.fill, cell.border = OK_FILL, BOX
ws.row_dimensions[r_mr + 1].height = 32

# --- mala tabela: broj i iznos po statusu (formule) ---
r0 = r_mr + 4
ws.cell(row=r0, column=2, value="Breakdown po statusu").font = SEC
hdr = ["Status", "Transactions", "% of Total", "Total Amount"]
for j, h in enumerate(hdr):
    c = ws.cell(row=r0 + 1, column=2 + j, value=h)
    c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BOX
    c.alignment = Alignment(horizontal="center")

for k, st in enumerate(["Matched", "Only in GL", "Only in Bank", "Amount Mismatch"]):
    r = r0 + 2 + k
    ws.cell(row=r, column=2, value=st).font = BODY
    ws.cell(row=r, column=3, value=f'=COUNTIF({R_STAT},$B{r})').number_format = "#,##0"
    ws.cell(row=r, column=4,
            value=f'=IFERROR($C{r}/COUNTA({R_ID}),0)').number_format = "0.0%"
    ws.cell(row=r, column=5,
            value=f'=SUMIF({R_STAT},$B{r},{R_AMT})').number_format = CUR
    for j in range(2, 6):
        cc = ws.cell(row=r, column=j)
        cc.border = BOX
        if j > 2:
            cc.font = BODY

r_tot = r0 + 6
ws.cell(row=r_tot, column=2, value="TOTAL").font = Font(name=FONT, bold=True, size=10)
ws.cell(row=r_tot, column=3, value=f"=SUM(C{r0+2}:C{r0+5})").number_format = "#,##0"
ws.cell(row=r_tot, column=4, value=f"=SUM(D{r0+2}:D{r0+5})").number_format = "0.0%"
ws.cell(row=r_tot, column=5, value=f"=SUM(E{r0+2}:E{r0+5})").number_format = CUR
for j in range(2, 6):
    cc = ws.cell(row=r_tot, column=j)
    cc.font = Font(name=FONT, bold=True, size=10)
    cc.fill, cc.border = PatternFill("solid", fgColor="DCE3EE"), BOX

# --- napomene ---
r_n = r_tot + 3
ws.cell(row=r_n, column=2, value="Napomene / pretpostavke").font = SEC
notes = [
    "Podaci su sinteticki (generisani u Pythonu, seed=42) — nisu stvarne transakcije.",
    "Merge je radjen po Transaction_ID (outer join). Prije merge-a je svaka tabela "
    "de-duplicirana (keep='first') da jedan ID daje jedan red.",
    "Amount Mismatch = |Amount_GL - Amount_Bank| > 0.01 (tolerancija na zaokruzivanje).",
    "'Duplicates' broji jedinstvene Transaction_ID-eve koji se pojavljuju vise od "
    "jednom u GL ili Bank tabeli.",
    "Bank Statement nema Category/Account — bank-only redovi su oznaceni kao "
    "'Unclassified (Bank only)'.",
    "Datumi u banci su 0-3 dana kasniji od GL (clearing delay); to nije tretirano "
    "kao izuzetak.",
]
for i, t in enumerate(notes):
    c = ws.cell(row=r_n + 1 + i, column=2, value="•  " + t)
    c.font = Font(name=FONT, size=9, color="404040")
    c.alignment = Alignment(wrap_text=True, vertical="top")

ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 26
for col in "CDEFGHI":
    ws.column_dimensions[col].width = 18

wb.save(OUT)
print("Sacuvano:", OUT)
